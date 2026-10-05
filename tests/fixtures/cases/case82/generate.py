"""Generate case82: Footnotes in Columns

Tests:
- Footnotes in a two-column section sit at the foot of the column that cites
  them, in the column's width; only that column gets shorter
- A three-column section (new page) with notes cited from columns 2 and 3
- No FootnoteText style (python-docx's template has none): notes take the
  default paragraph style and the document's spacing defaults
python-docx has no footnote API, so footnotes.xml is written by hand.
"""

import io
import pathlib
import random
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

WML = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORDS = (
    "lorem ipsum dolor sit amet consectetur adipiscing elit sed do eiusmod tempor "
    "incididunt ut labore et dolore magna aliqua"
).split()
rnd = random.Random(82)


def words(k):
    return " ".join(rnd.choice(WORDS) for _ in range(k))


def set_cols(section, num):
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:cols")):
        sect_pr.remove(old)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), "720")
    grid = sect_pr.find(qn("w:docGrid"))
    if grid is not None:
        grid.addprevious(cols)
    else:
        sect_pr.append(cols)


notes = []


def cite(paragraph):
    notes.append(f"Footnote {len(notes) + 1}: " + words(25) + ".")
    run = paragraph.add_run()
    rpr = OxmlElement("w:rPr")
    va = OxmlElement("w:vertAlign")
    va.set(qn("w:val"), "superscript")
    rpr.append(va)
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), str(len(notes)))
    run._r.append(rpr)
    run._r.append(ref)


doc = Document()
for i in range(14):
    p = doc.add_paragraph(f"P{i} " + words(45))
    if i in (1, 9):
        cite(p)
set_cols(doc.sections[-1], 2)

doc.add_section(WD_SECTION.NEW_PAGE)
for i in range(12):
    p = doc.add_paragraph(f"Q{i} " + words(30))
    if i in (5, 9):
        cite(p)
set_cols(doc.sections[-1], 3)

buf = io.BytesIO()
doc.save(buf)
buf.seek(0)

note_xml = "".join(
    f'<w:footnote w:id="{k + 1}"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
    '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
    f'<w:r><w:t xml:space="preserve"> {text}</w:t></w:r></w:p></w:footnote>'
    for k, text in enumerate(notes)
)
footnotes = (
    f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{WML}">'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    + note_xml
    + "</w:footnotes>"
)

out_buf = io.BytesIO()
with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "[Content_Types].xml":
            data = data.replace(
                b"</Types>",
                b'<Override PartName="/word/footnotes.xml" ContentType="application/'
                b'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
            )
        elif item.filename == "word/_rels/document.xml.rels":
            data = data.replace(
                b"</Relationships>",
                b'<Relationship Id="rIdFootnotes" Type="http://schemas.openxmlformats.org/'
                b'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>',
            )
        elif item.filename == "word/settings.xml":
            tree = etree.fromstring(data)
            for compat_setting in tree.iter("{%s}compatSetting" % WML):
                if compat_setting.get(qn("w:name")) == "compatibilityMode":
                    compat_setting.set(qn("w:val"), "15")
            data = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        zout.writestr(item, data)
    zout.writestr("word/footnotes.xml", footnotes)

out_path = pathlib.Path(__file__).parent / "input.docx"
out_path.write_bytes(out_buf.getvalue())
print(f"Wrote {out_path} ({out_path.stat().st_size} bytes)")
