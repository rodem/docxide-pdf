"""Generate case79: where Word puts text inside docGrid type="lines" cells.

Every paragraph is one line with no spacing before or after and single line
spacing, so each line starts where the previous one's cells end and its
baseline offset inside the cell can be read straight off the PDF. Each line
uses one font for all its characters (ascii, hAnsi and eastAsia set alike) and
starts with a label naming the font and size.

Section 1 has an 18pt pitch, section 2 (new page) a 15.6pt pitch, the two
pitches Word's Western and Chinese templates use. Sizes are picked so some
lines need two cells (20pt Times New Roman, 16pt YaHei on 15.6pt), and the
Latin fonts span a wide range of line gaps (Times New Roman 0.04 em, Calibri
0, Yu Mincho 0.5). The last lines of section 1 test a run mix (Times New Roman
beside MS Mincho) and 1.5-line spacing on the grid.
"""
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

JA = "日本語のテキスト"
ZH = "中文文本测试"
TW = "繁體中文測試"
KO = "한국어 텍스트"

SECTION_1 = [
    ("Times New Roman", 8, None),
    ("Times New Roman", 10, None),
    ("Times New Roman", 12, None),
    ("Times New Roman", 14, None),
    ("Times New Roman", 20, None),
    ("Arial", 10, None),
    ("Arial", 12, None),
    ("Calibri", 11, None),
    ("Calibri", 16, None),
    ("Aptos", 12, None),
    ("Cambria", 12, None),
    ("Georgia", 12, None),
    ("Courier New", 12, None),
    ("Verdana", 10, None),
    ("MS Mincho", 10.5, JA),
    ("MS Mincho", 16, JA),
    ("MS Gothic", 12, JA),
    ("Yu Mincho", 10.5, JA),
    ("Yu Mincho", 14, JA),
    ("SimSun", 10.5, ZH),
    ("SimSun", 12, ZH),
    ("Microsoft YaHei", 12, ZH),
    ("Microsoft YaHei", 16, ZH),
    ("PMingLiU", 12, TW),
    ("Malgun Gothic", 11, KO),
]

SECTION_2 = [
    ("Times New Roman", 12, None),
    ("Calibri", 11, None),
    ("Arial", 10, None),
    ("MS Mincho", 10.5, JA),
    ("SimSun", 12, ZH),
    ("Microsoft YaHei", 12, ZH),
    ("Microsoft YaHei", 16, ZH),
    ("Yu Mincho", 10.5, JA),
]


def set_grid(section, pitch_twips: int) -> None:
    sect_pr = section._sectPr
    grid = sect_pr.find(qn("w:docGrid"))
    if grid is None:
        grid = OxmlElement("w:docGrid")
        sect_pr.append(grid)
    grid.set(qn("w:type"), "lines")
    grid.set(qn("w:linePitch"), str(pitch_twips))


def tight(paragraph, line_spacing: float = 1.0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = line_spacing


def add_run(paragraph, text: str, font: str, size: float) -> None:
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), font)


def add_line(doc, font: str, size: float, sample, line_spacing: float = 1.0) -> None:
    p = doc.add_paragraph()
    tight(p, line_spacing)
    label = f"{font} {size:g}pt" + (f" x{line_spacing:g}" if line_spacing != 1.0 else "")
    add_run(p, f"{label} {sample or 'Hxgp'}", font, size)


doc = Document()
for style_name in ("Normal",):
    tight(doc.styles[style_name])
set_grid(doc.sections[0], 360)

for font, size, sample in SECTION_1:
    add_line(doc, font, size, sample)

mixed = doc.add_paragraph()
tight(mixed)
add_run(mixed, "Times New Roman 12pt beside ", "Times New Roman", 12)
add_run(mixed, f"MS Mincho 12pt {JA}", "MS Mincho", 12)
add_line(doc, "Times New Roman", 12, None, 1.5)
add_line(doc, "MS Mincho", 10.5, JA, 1.5)

set_grid(doc.add_section(WD_SECTION.NEW_PAGE), 312)
for font, size, sample in SECTION_2:
    add_line(doc, font, size, sample)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "input.docx")
doc.save(out)
