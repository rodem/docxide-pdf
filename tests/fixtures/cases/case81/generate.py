"""Generate case81: Column Balancing

Every column region here ends at a continuous section break, which makes Word
balance it. Tests:
- A three-column region spanning pages: only its last page is balanced, and
  Word's split there (14/14/12 lines) is not the shortest one that fits
- One long paragraph balanced over two columns (the region ends at its
  content height over the column count, above the paragraph's space after)
- Double line spacing over three columns
- Headings with space before at a region start and inside columns
- Short paragraphs that cannot split (widow control), leaving the third
  column empty
- A region with a column break: not balanced, column 2 runs to the page foot
"""

import io
import pathlib
import random
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

WML = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORDS = (
    "lorem ipsum dolor sit amet consectetur adipiscing elit sed do eiusmod tempor "
    "incididunt ut labore et dolore magna aliqua"
).split()
rnd = random.Random(81)


def words(k):
    return " ".join(rnd.choice(WORDS) for _ in range(k))


def set_cols(section, num, sep=False):
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:cols")):
        sect_pr.remove(old)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), "720")
    if sep:
        cols.set(qn("w:sep"), "1")
    grid = sect_pr.find(qn("w:docGrid"))
    if grid is not None:
        grid.addprevious(cols)
    else:
        sect_pr.append(cols)


def one_column(text):
    doc.add_section(WD_SECTION.CONTINUOUS)
    doc.add_paragraph(text)
    set_cols(doc.sections[-1], 1)


doc = Document()
doc.add_heading("Column Balancing", level=1)
set_cols(doc.sections[-1], 1)

# Three columns over more than a page
doc.add_section(WD_SECTION.CONTINUOUS)
for i in range(50):
    doc.add_paragraph(f"A{i} Lorem ipsum dolor sit amet, consectetur adipiscing elit. "
                      "Sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. "
                      "Ut enim ad minim veniam.")
set_cols(doc.sections[-1], 3, sep=True)

one_column("One long paragraph over two columns follows.")

doc.add_section(WD_SECTION.CONTINUOUS)
doc.add_paragraph("B " + words(260) + ".")
set_cols(doc.sections[-1], 2, sep=True)

one_column("Double line spacing over three columns follows.")

doc.add_section(WD_SECTION.CONTINUOUS)
for i, k in enumerate([50, 70, 40, 60]):
    p = doc.add_paragraph(f"C{i} " + words(k) + ".")
    p.paragraph_format.line_spacing = 2.0
set_cols(doc.sections[-1], 3)

one_column("Headings in two columns follow.")

doc.add_section(WD_SECTION.CONTINUOUS)
for i in range(4):
    doc.add_heading(f"Heading D{i}", level=2)
    doc.add_paragraph(f"D{i} " + words(45 + 10 * i) + ".")
set_cols(doc.sections[-1], 2, sep=True)

one_column("Four short paragraphs in three columns follow.")

doc.add_section(WD_SECTION.CONTINUOUS)
for i in range(4):
    doc.add_paragraph(f"E{i} " + words(10) + ".")
set_cols(doc.sections[-1], 3, sep=True)

one_column("A region with a column break follows.")

doc.add_section(WD_SECTION.CONTINUOUS)
doc.add_paragraph("F0 " + words(40) + ".")
doc.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
doc.add_paragraph("F1 " + words(30) + ".")
doc.add_paragraph("F2 " + words(120) + ".")
set_cols(doc.sections[-1], 4, sep=True)

one_column("End of document.")

tmp_buf = io.BytesIO()
doc.save(tmp_buf)
tmp_buf.seek(0)

out_buf = io.BytesIO()
with zipfile.ZipFile(tmp_buf, "r") as zin, zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "word/settings.xml":
            tree = etree.fromstring(data)
            for compat_setting in tree.iter("{%s}compatSetting" % WML):
                if compat_setting.get(qn("w:name")) == "compatibilityMode":
                    compat_setting.set(qn("w:val"), "15")
            data = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        zout.writestr(item, data)

out_path = pathlib.Path(__file__).parent / "input.docx"
out_path.write_bytes(out_buf.getvalue())
print(f"Wrote {out_path} ({out_path.stat().st_size} bytes)")
